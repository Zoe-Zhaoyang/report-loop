import importlib.util
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

spec = importlib.util.spec_from_file_location("workflow", Path(__file__).resolve().parents[1] / "scripts" / "structured_data.py")
w = importlib.util.module_from_spec(spec)
spec.loader.exec_module(w)


def evidence(case="中文 case"):
    return {"schema": "openharness-structured-data/v1", "case_id": case,
            "items": [{"id": "EV-001", "type": "quantitative", "source_ref": "访谈.txt", "content": "样本12人"}], "unresolved": []}


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="中文 路径 (test)-")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / "素材"
        self.source.mkdir()
        (self.source / "访谈.txt").write_text("样本12人，其中8人每周使用3次", encoding="utf-8")

    def test_schema_rejects_duplicate_ids_wrong_case_and_extra_fields(self):
        w.validate(evidence(), "中文 case")
        bad = evidence()
        bad["items"].append(bad["items"][0].copy())
        with self.assertRaises(w.WorkflowError): w.validate(bad, "中文 case")
        with self.assertRaises(w.WorkflowError): w.validate(evidence(), "other")
        bad = evidence(); bad["extra"] = 1
        with self.assertRaises(w.WorkflowError): w.validate(bad, "中文 case")

    def test_collect_unicode_and_unsupported(self):
        (self.source / "截图.png").write_bytes(b"not a real image")
        data, inventory, unresolved = w.collect([str(self.source)], self.root / "out", 1000)
        self.assertIn("样本12人", data[0]["text"])
        self.assertEqual(len(inventory), 2)
        self.assertTrue(any("截图.png" in x for x in unresolved))

    def test_output_overlap_and_size_limit(self):
        with self.assertRaises(w.WorkflowError): w.collect([str(self.source)], self.source / "out", 1000)
        with self.assertRaises(w.WorkflowError): w.collect([str(self.source)], self.root, 1000)
        with self.assertRaises(w.WorkflowError): w.collect([str(self.source)], self.root / "out", 2)

    def test_gb18030_and_csv_multiline(self):
        file = self.source / "数据.csv"
        file.write_bytes('名称,备注\n甲,"第一行\n第二行"\n'.encode("gb18030"))
        rows, _ = w.extract(file)
        self.assertIn("row 2", rows[0][1])
        self.assertNotIn("row 3", rows[0][1])

    def test_docx_order_and_xlsx_missing_formula_cache(self):
        try:
            from docx import Document
            from openpyxl import Workbook
        except ImportError:
            self.skipTest("optional dependencies unavailable")
        doc = Document(); doc.add_paragraph("前文")
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "表格"
        doc.add_paragraph("后文"); path = self.source / "访谈.docx"; doc.save(path)
        sections, warnings = w.extract(path)
        self.assertEqual([v for _, v in sections], ["前文", '["表格"]', "后文"])
        book = Workbook(); book.active["A1"] = "=1+2"
        book.create_sheet("隐藏").sheet_state = "hidden"
        book["隐藏"]["A1"] = "仍需读取"
        path = self.source / "指标.xlsx"; book.save(path)
        sections, warnings = w.extract(path)
        self.assertTrue(any("没有缓存" in x for x in warnings))
        self.assertTrue(any("仍需读取" in x for _, x in sections))

    def test_pdf_blank_page_unresolved(self):
        try: from pypdf import PdfWriter
        except ImportError: self.skipTest("pypdf unavailable")
        path = self.source / "扫描件.pdf"
        writer = PdfWriter(); writer.add_blank_page(width=100, height=100)
        with path.open("wb") as f: writer.write(f)
        sections, warnings = w.extract(path)
        self.assertFalse(sections)
        self.assertTrue(any("OCR" in x for x in warnings))

    def test_real_subprocess_unicode_transport(self):
        fake = self.root / "模拟 cli.py"
        fake.write_text('import sys,json\nfrom pathlib import Path\nprompt=sys.stdin.buffer.read().decode("utf-8")\nPath(sys.argv[1]).write_text(json.dumps({"prompt":prompt},ensure_ascii=False),encoding="utf-8")\n', encoding="utf-8")
        result = self.root / "结果.json"
        payload = w.invoke([sys.executable, str(fake), str(result)], "中文 & 空格 (你好)", result, self.root, 10)
        self.assertEqual(payload["prompt"], "中文 & 空格 (你好)")

    def test_failure_timeout_and_invalid_json(self):
        result = self.root / "missing.json"
        with self.assertRaises(w.WorkflowError):
            w.invoke([sys.executable, "-c", "raise SystemExit(7)"], "", result, self.root, 10)
        with self.assertRaises(w.WorkflowError):
            w.invoke([sys.executable, "-c", "import time; time.sleep(30)"], "", result, self.root, 0.1)
        result.write_text("not json", encoding="utf-8")
        with self.assertRaises(ValueError):
            w.invoke([sys.executable, "-c", "pass"], "", result, self.root, 10)

    def test_full_flow_and_no_overwrite(self):
        output = self.root / "输出"
        args = ["run", "--source", str(self.source), "--case-id", "中文 case", "--output", str(output)]
        with patch.object(w, "codex_command", return_value=["unused"]), patch.object(w, "invoke", return_value=evidence()):
            self.assertEqual(w.main(args), 0)
            self.assertEqual(w.main(args), 2)
        result = json.loads((output / "structured_data.json").read_text(encoding="utf-8"))
        self.assertEqual(result["case_id"], "中文 case")
        self.assertTrue((output / "run_manifest.json").is_file())

    def test_model_failure_no_output(self):
        output = self.root / "失败输出"
        with patch.object(w, "codex_command", return_value=["unused"]), patch.object(w, "invoke", side_effect=w.WorkflowError("failed")):
            self.assertEqual(w.main(["run", "--source", str(self.source), "--case-id", "x", "--output", str(output)]), 2)
        self.assertFalse(output.exists())

    def test_npm_shim_resolution(self):
        shim = self.root / "codex.cmd"; shim.write_text("@echo off", encoding="utf-8")
        entry = self.root / "node_modules/@openai/codex/bin/codex.js"
        entry.parent.mkdir(parents=True); entry.write_text("", encoding="utf-8")
        with patch.object(w.shutil, "which", return_value="node.exe"):
            self.assertEqual(w.codex_command(str(shim)), ["node.exe", str(entry.resolve())])


if __name__ == "__main__": unittest.main()
